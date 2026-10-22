"""Source-to-Markdown conversion service.

Converts any source document (PDF/Excel/Word/PPT/TXT) to Obsidian-style
Markdown, then optimizes with qwen-max: generate summary, insert wikilinks
to existing vault notes, extract key concepts as tags.

Inspired by claude-obsidian wiki-ingest pattern: capture source, classify,
extract structured content, LLM-enhance, save to vault with provenance.
"""
import logging

logger = logging.getLogger(__name__)

import os
import re
import json
import asyncio
import httpx
import requests
from datetime import datetime
from typing import Callable, List, Optional

from app.config import settings
from app.services.model_gateway import endpoint, headers as gateway_headers
from app.services.knowledge.progress import KnowledgeImportCancelled
from app.services.knowledge.doc_images import expand_placeholders, make_image_sink

COMPAT_URL = endpoint("chat/completions")
# 单次改写超过该字符数时按段落边界分段顺序改写（全文覆盖，不截断、不丢内容）
REWRITE_PART_CHARS = 60000

ProgressCallback = Callable[[str, int, str], None]
CancelCheck = Callable[[], bool]


def _looks_binary(data: bytes) -> bool:
    """粗判内容是否为二进制（PDF/Office/压缩包等）：头部含 NUL 或已知文件魔数。"""
    head = data[:8192]
    if b"\x00" in head:
        return True
    sig = data[:8]
    return (
        sig.startswith(b"%PDF")          # PDF
        or sig.startswith(b"PK\x03\x04")  # zip 系（docx/xlsx/pptx）
        or sig.startswith(b"\xd0\xcf\x11\xe0")  # OLE2（doc/xls/ppt）
    )


def _llm(system: str, user: str, model: str = None, temperature: float = 0.3) -> str:
    headers = gateway_headers()
    payload = {
        "model": model or settings.KNOWLEDGE_REWRITE_MODEL,
        "messages": [
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
        "temperature": temperature,
    }
    r = requests.post(COMPAT_URL, headers=headers, json=payload, timeout=600)
    r.raise_for_status()
    return r.json()["choices"][0]["message"]["content"]


async def _llm_async(
    system: str,
    user: str,
    model: str = None,
    temperature: float = 0.3,
    on_chunk=None,
) -> str:
    headers = gateway_headers()
    payload = {
        "model": model or settings.KNOWLEDGE_REWRITE_MODEL,
        "messages": [
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
        "temperature": temperature,
    }
    async with httpx.AsyncClient(timeout=600.0) as client:
        if on_chunk is None:
            response = await client.post(COMPAT_URL, headers=headers, json=payload)
            response.raise_for_status()
            return response.json()["choices"][0]["message"]["content"]

        payload["stream"] = True
        pieces = []
        async with client.stream("POST", COMPAT_URL, headers=headers, json=payload) as response:
            response.raise_for_status()
            async for line in response.aiter_lines():
                if not line.startswith("data:"):
                    continue
                data = line[5:].strip()
                if data == "[DONE]":
                    break
                try:
                    chunk = json.loads(data)
                except json.JSONDecodeError:
                    continue
                choices = chunk.get("choices") or []
                if not choices:
                    continue
                delta = choices[0].get("delta") or {}
                text = delta.get("content")
                if text:
                    pieces.append(text)
                    on_chunk(text)
        return "".join(pieces)


class SourceToMdService:
    """Convert source documents to Obsidian-style Markdown with LLM optimization."""

    def convert_to_md(
        self,
        source_path: str,
        vault_dir: str,
        original_filename: str,
        stored_name: str,
        existing_titles: Optional[List[str]] = None,
        kb_id: int = 0,
    ) -> str:
        """Main entry: source file to MD file path.

        Returns the absolute path of the generated .md file in the vault.
        """
        ext = os.path.splitext(original_filename)[1].lower()
        self._check_extractor_available(ext)
        sink = make_image_sink(vault_dir)
        raw_md = self._extract_to_markdown(source_path, original_filename, ext, sink=sink)
        if not raw_md.strip():
            raise ValueError("Document extraction yielded no content")

        # 规则清洗：去页眉页脚/页码/作者单位/致谢/参考文献等噪声。
        # 清洗后的原文用于向量入库，同时作为 LLM 改写的输入。
        from app.services.knowledge.raw_cleaner import clean_extracted_text
        cleaned = clean_extracted_text(raw_md) or raw_md

        # LLM 全文改写：summary + wikilinks + tags。失败时降级为清洗原文，
        # 不阻塞 md 生成与向量化。
        titles = existing_titles or []
        try:
            optimized = self._llm_optimize(cleaned, titles, original_filename)
        except RuntimeError as e:
            logger.info(f"[source_to_md] LLM rewrite failed, degrade to cleaned raw text: {e}")
            optimized = {"title": "", "summary": "", "tags": [], "body": cleaned}

        # Build final MD with frontmatter + source link
        title = optimized.get("title") or os.path.splitext(original_filename)[0]
        tags = optimized.get("tags", [])
        summary = optimized.get("summary", "")
        enhanced_body = optimized.get("body") or cleaned
        # 图片占位符展开：md 里是 ![图像]+描述块（阅读可见）
        enhanced_body = expand_placeholders(enhanced_body, sink.images if sink else {}, for_raw=False)

        frontmatter = self._build_frontmatter(
            title=title,
            source_file=stored_name,
            original_filename=original_filename,
            tags=tags,
            summary=summary,
        )

        source_link = self._source_link(stored_name, ext, kb_id)
        full_md = frontmatter + enhanced_body + source_link

        # 旁挂原文：占位符展开为纯文本块（进入切块/向量化/FTS，可被检索命中）
        cleaned_raw = expand_placeholders(cleaned, sink.images if sink else {}, for_raw=True)
        return self._save_markdown(vault_dir, original_filename, full_md, cleaned_raw)

    async def convert_to_md_async(
        self,
        source_path: str,
        vault_dir: str,
        original_filename: str,
        stored_name: str,
        existing_titles: Optional[List[str]] = None,
        kb_id: int = 0,
        on_progress: Optional[ProgressCallback] = None,
        should_cancel: Optional[CancelCheck] = None,
    ) -> str:
        """Async conversion path used by upload ingestion.

        The async LLM call allows a cancellation request to stop a long rewrite
        instead of waiting for the full HTTP timeout.
        """
        ext = os.path.splitext(original_filename)[1].lower()
        if on_progress:
            on_progress("抽取正文", 12, "读取源文件")
        if should_cancel and should_cancel():
            raise KnowledgeImportCancelled()
        self._check_extractor_available(ext)

        def _img_progress(n: int) -> None:
            if on_progress:
                on_progress("理解图片", 16, f"已理解 {n} 张文档图片")

        sink = make_image_sink(vault_dir, on_image=_img_progress if on_progress else None)
        raw_md = await asyncio.to_thread(
            self._extract_to_markdown, source_path, original_filename, ext, sink
        )
        if should_cancel and should_cancel():
            raise KnowledgeImportCancelled()
        if not raw_md.strip():
            raise ValueError("Document extraction yielded no content")
        if on_progress:
            on_progress("清洗正文", 24, "清理页眉页脚等噪声")

        from app.services.knowledge.raw_cleaner import clean_extracted_text
        cleaned = clean_extracted_text(raw_md) or raw_md
        if on_progress:
            on_progress("智能改写", 32, "准备分段改写")

        titles = existing_titles or []
        try:
            optimized = await self._llm_optimize_async(
                cleaned,
                titles,
                original_filename,
                on_progress=on_progress,
                should_cancel=should_cancel,
            )
        except RuntimeError as e:
            logger.info(f"[source_to_md] LLM rewrite failed, degrade to cleaned raw text: {e}")
            optimized = {"title": "", "summary": "", "tags": [], "body": cleaned}

        if should_cancel and should_cancel():
            raise KnowledgeImportCancelled()
        if on_progress:
            on_progress("生成笔记", 72, "写入 Markdown")

        title = optimized.get("title") or os.path.splitext(original_filename)[0]
        tags = optimized.get("tags", [])
        summary = optimized.get("summary", "")
        enhanced_body = optimized.get("body") or cleaned
        # 图片占位符展开：md 里是 ![图像]+描述块（阅读可见）
        enhanced_body = expand_placeholders(
            enhanced_body, sink.images if sink else {}, for_raw=False
        )
        full_md = (
            self._build_frontmatter(
                title=title,
                source_file=stored_name,
                original_filename=original_filename,
                tags=tags,
                summary=summary,
            )
            + enhanced_body
            + self._source_link(stored_name, ext, kb_id)
        )
        cleaned_raw = expand_placeholders(cleaned, sink.images if sink else {}, for_raw=True)
        return self._save_markdown(vault_dir, original_filename, full_md, cleaned_raw)

    @staticmethod
    def _save_markdown(vault_dir: str, original_filename: str, full_md: str, cleaned: str) -> str:
        md_filename = os.path.splitext(original_filename)[0]
        md_filename = re.sub(r'[\\/:*?"<>|]', "_", md_filename).strip() or "converted"
        md_filename += ".md"
        md_path = os.path.join(vault_dir, md_filename)

        if os.path.exists(md_path):
            base, suffix = os.path.splitext(md_filename)
            index = 1
            while os.path.exists(os.path.join(vault_dir, f"{base}_{index}{suffix}")):
                index += 1
            md_path = os.path.join(vault_dir, f"{base}_{index}{suffix}")

        os.makedirs(vault_dir, exist_ok=True)
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(full_md)

        # 旁挂清洗后的原文，向量化切块优先读取它；非 .md 后缀不进图谱。
        raw_path = os.path.splitext(md_path)[0] + ".raw.txt"
        with open(raw_path, "w", encoding="utf-8") as f:
            f.write(cleaned)
        return md_path

    # ---- Format-specific extraction ----

    @staticmethod
    def _check_extractor_available(ext: str) -> None:
        """按扩展名预检解析依赖：缺失时快速失败并给出明确指引。

        不做预检时各 _xxx_to_md 会静默降级到 _fallback_text，把二进制当文本读，
        产生乱码进 LLM 与向量库（服务器缺依赖时的典型故障）。
        """
        import importlib
        required = {
            ".pdf": ("fitz", "pymupdf"),
            ".xlsx": ("openpyxl", "openpyxl"),
            ".xls": ("openpyxl", "openpyxl"),
            ".doc": ("docx", "python-docx"),
            ".docx": ("docx", "python-docx"),
            ".ppt": ("pptx", "python-pptx"),
            ".pptx": ("pptx", "python-pptx"),
        }
        req = required.get(ext)
        if not req:
            return
        module_name, package_name = req
        try:
            importlib.import_module(module_name)
        except ImportError:
            raise ValueError(
                f"缺少 {ext} 文件的解析依赖 {package_name}，无法抽取正文。"
                f"请在部署环境安装后重试：pip install {package_name}"
            )

    def _extract_to_markdown(self, path: str, filename: str, ext: str, sink=None) -> str:
        if ext == ".pdf":
            # OCR 前置层：扫描版 PDF（文字层稀疏）先本地 RapidOCR 抢救正文，
            # 有结果则直接用（跳过字号分析——OCR 文本没有字号信息）
            try:
                from app.config import settings
                if getattr(settings, "PDF_OCR_ENABLED", False):
                    from app.services.knowledge.pdf_ocr import (
                        extract_pdf_with_ocr, is_scanned_pdf, ocr_available,
                    )
                    if ocr_available():
                        scanned, pages, chars = is_scanned_pdf(path)
                        if scanned:
                            logger.info(
                                f"[ocr前置] 检测到扫描版 PDF（{pages}页仅{chars}字符），走 OCR 提取: {filename}"
                            )
                            ocr_text = extract_pdf_with_ocr(path)
                            if ocr_text:
                                return ocr_text
                            logger.warning(f"[ocr前置] OCR 无结果，回退默认提取: {filename}")
            except Exception as e:
                logger.warning(f"[ocr前置] 处理失败，回退默认提取: {e}")
            return self._pdf_to_md(path, sink=sink)
        if ext in (".xlsx", ".xls", ".csv"):
            return self._excel_to_md(path, ext, sink=sink)
        if ext in (".doc", ".docx"):
            return self._doc_to_md(path, sink=sink)
        if ext in (".ppt", ".pptx"):
            return self._ppt_to_md(path, sink=sink)
        # Fallback: plain text
        try:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception:
            return ""

    def _pdf_to_md(self, pdf_path: str, sink=None) -> str:
        """Extract PDF text via PyMuPDF, format as Markdown with heading detection."""
        try:
            import fitz
        except ImportError:
            return self._fallback_text(pdf_path)

        doc = fitz.open(pdf_path)
        if len(doc) == 0:
            doc.close()
            return ""

        # Analyze font sizes to detect headings
        from collections import Counter
        size_counter = Counter()
        for page in doc:
            blocks = page.get_text("dict")["blocks"]
            for block in blocks:
                if block["type"] == 0:
                    for line in block["lines"]:
                        for span in line["spans"]:
                            text = span["text"].strip()
                            if text:
                                size_counter[round(span["size"], 1)] += len(text)

        body_size = 12
        heading_sizes = []
        if size_counter:
            sorted_sizes = sorted(size_counter.items(), key=lambda x: x[1], reverse=True)
            body_size = sorted_sizes[0][0]
            all_sizes = sorted(size_counter.keys(), reverse=True)
            heading_sizes = [s for s in all_sizes if s > body_size + 1][:3]

        def get_heading_level(size):
            for i, hs in enumerate(heading_sizes):
                if size >= hs - 0.5:
                    return min(i + 1, 3)
            return 0

        # Detect repeated headers/footers
        noise_texts = set()
        if len(doc) >= 3:
            scan_pages = list(range(min(len(doc), 20)))
            if len(doc) > 20:
                scan_pages += list(range(max(0, len(doc) - 20), len(doc)))
            header_counter = Counter()
            footer_counter = Counter()
            for i in scan_pages:
                page = doc[i]
                rect = page.rect
                h = rect.height
                top_rect = fitz.Rect(0, 0, rect.width, h * 0.15)
                bottom_rect = fitz.Rect(0, h * 0.85, rect.width, h)
                for b in page.get_text("blocks"):
                    b_rect = fitz.Rect(b[:4])
                    text = b[4].strip()
                    if not text:
                        continue
                    if b_rect.intersects(top_rect):
                        header_counter[text] += 1
                    elif b_rect.intersects(bottom_rect):
                        footer_counter[text] += 1
            for counter in [header_counter, footer_counter]:
                for text, count in counter.items():
                    if count / len(scan_pages) > 0.6:
                        noise_texts.add(text)

        parts = []
        for page_num, page in enumerate(doc, 1):
            if page_num > 1:
                parts.append(f"\n\n<!-- Page {page_num} -->\n")
            blocks = page.get_text("dict")["blocks"]
            for block in blocks:
                if block["type"] != 0:
                    # 图片块：存档 + VLM 理解，原位置插占位符（不可用时照旧跳过）
                    if sink is not None:
                        ph = sink.add(
                            block.get("image"),
                            block.get("ext") or "png",
                            bbox=block.get("bbox"),
                        )
                        if ph:
                            parts.append(ph)
                    continue
                for line in block["lines"]:
                    spans = line["spans"]
                    if not spans:
                        continue
                    line_text = "".join(s["text"] for s in spans).strip()
                    if not line_text or line_text in noise_texts:
                        continue
                    max_size = max(s["size"] for s in spans)
                    level = get_heading_level(max_size)
                    if level > 0 and len(line_text) < 80:
                        clean = re.sub(r"\*+", "", line_text)
                        parts.append(f"\n{'#' * level} {clean}\n")
                    else:
                        parts.append(line_text + "\n")
        doc.close()
        md = "\n".join(parts)
        md = re.sub(r"\n{3,}", "\n\n", md).strip()
        return md

    def _excel_to_md(self, path: str, ext: str, sink=None) -> str:
        """Extract Excel/CSV to Markdown tables."""
        try:
            import openpyxl
        except ImportError:
            return self._fallback_text(path)

        img_placeholders: List[str] = []
        if ext == ".csv":
            import csv
            rows = []
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                for row in reader:
                    rows.append(row)
        else:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                if ws.title and len(wb.worksheets) > 1:
                    rows.append([f"## {ws.title}"])
                for row in ws.iter_rows(values_only=True):
                    if any(c is not None for c in row):
                        rows.append([str(c) if c is not None else "" for c in row])
                # 浮动图片（私有 API，尽力而为；失败静默跳过）
                if sink is not None:
                    for im in getattr(ws, "_images", None) or []:
                        try:
                            ph = sink.add(im._data(), getattr(im, "format", None) or "png")
                            if ph:
                                img_placeholders.append(ph)
                        except Exception:
                            pass
            wb.close()

        if not rows:
            return ""

        parts = []
        for row in rows:
            if len(row) == 1 and isinstance(row[0], str) and row[0].startswith("## "):
                parts.append(f"\n{row[0]}\n")
                continue
            cells = [str(c).replace("|", "\\|").replace("\n", " ")[:100] for c in row]
            parts.append("| " + " | ".join(cells) + " |")
            if len(parts) == 1 or (len(parts) >= 2 and parts[-2].startswith("|") and not parts[-2].startswith("|---")):
                # Add header separator after first data row of a table
                pass

        # Build proper markdown tables
        md_lines = []
        i = 0
        while i < len(parts):
            line = parts[i]
            if line.startswith("## "):
                md_lines.append(line)
                i += 1
                continue
            if line.startswith("|"):
                md_lines.append(line)
                # Check if next is also a table row
                if i + 1 < len(parts) and parts[i + 1].startswith("|") and not parts[i + 1].startswith("|---"):
                    col_count = line.count("|") - 1
                    md_lines.append("| " + " | ".join(["---"] * col_count) + " |")
                i += 1
                continue
            md_lines.append(line)
            i += 1

        md = "\n".join(md_lines)
        # 浮动图片占位符追加在末尾（不能进表格行处理，会破坏表格格式化）
        if img_placeholders:
            md = md.rstrip("\n") + "\n" + "\n".join(img_placeholders) + "\n"
        return md

    def _doc_to_md(self, path: str, sink=None) -> str:
        """Extract Word document to Markdown."""
        try:
            from docx import Document
        except ImportError:
            return self._fallback_text(path)

        doc = Document(path)
        parts = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = ((para.style.name if para.style else "") or "").lower()
            if "heading 1" in style:
                parts.append(f"\n# {text}\n")
            elif "heading 2" in style:
                parts.append(f"\n## {text}\n")
            elif "heading 3" in style:
                parts.append(f"\n### {text}\n")
            elif "list" in style:
                parts.append(f"- {text}")
            else:
                parts.append(text + "\n")
            # 段落内嵌图片：存档 + 理解，插在所属段落之后
            if sink is not None:
                try:
                    from docx.oxml.ns import qn
                    for blip in para._p.findall(".//" + qn("a:blip")):
                        rid = blip.get(qn("r:embed"))
                        if not rid:
                            continue
                        part = doc.part.related_parts.get(rid)
                        if part is None:
                            continue
                        ph = sink.add(
                            part.blob,
                            os.path.splitext(str(part.partname))[1].lstrip(".") or "png",
                        )
                        if ph:
                            parts.append(ph)
                except Exception as e:
                    logger.warning(f"[source_to_md] docx 图片抽取失败: {e}")

        for table in doc.tables:
            parts.append("")
            for row in table.rows:
                cells = [c.text.strip().replace("|", "\\|")[:100] for c in row.cells]
                parts.append("| " + " | ".join(cells) + " |")
            if table.rows:
                col_count = len(table.rows[0].cells)
                parts.insert(-len(table.rows), "| " + " | ".join(["---"] * col_count) + " |")
            parts.append("")

        return "\n".join(parts)

    def _ppt_to_md(self, path: str, sink=None) -> str:
        """Extract PowerPoint to Markdown."""
        try:
            from pptx import Presentation
        except ImportError:
            return self._fallback_text(path)

        prs = Presentation(path)
        parts = []
        for slide_num, slide in enumerate(prs.slides, 1):
            parts.append(f"\n## Slide {slide_num}\n")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            if para.level > 0:
                                parts.append(f"  {'- ' * para.level}{text}")
                            else:
                                parts.append(text)
                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        cells = [c.text.strip().replace("|", "\\|")[:100] for c in row.cells]
                        parts.append("| " + " | ".join(cells) + " |")
                    if table.rows:
                        col_count = len(table.rows[0].cells)
                        parts.insert(-len(table.rows), "| " + " | ".join(["---"] * col_count) + " |")
                # 幻灯片图片：存档 + 理解，插在该 shape 之后
                if sink is not None:
                    try:
                        from pptx.enum.shapes import MSO_SHAPE_TYPE
                        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                            img = shape.image
                            ph = sink.add(img.blob, img.ext or "png")
                            if ph:
                                parts.append(ph)
                    except Exception as e:
                        logger.warning(f"[source_to_md] pptx 图片抽取失败: {e}")
        return "\n".join(parts)

    def _fallback_text(self, path: str) -> str:
        """按纯文本读取；检测到二进制内容（缺解析库时的 PDF/Office 原文）时
        明确报错，而不是把二进制当文本喂给 LLM 与向量库。"""
        try:
            with open(path, "rb") as f:
                data = f.read()
        except Exception:
            return ""
        if _looks_binary(data):
            raise ValueError(
                "无法从该文件提取文本：文件为二进制格式且对应的解析依赖缺失"
                "（请检查部署环境是否安装了 pymupdf/openpyxl/python-docx/python-pptx）"
            )
        return data.decode("utf-8", errors="ignore")

    # ---- LLM optimization ----

    # ---- LLM optimization (full text, split into parts when needed) ----

    @staticmethod
    def _split_for_rewrite(text: str) -> List[str]:
        """按段落边界把全文拆成若干段顺序改写。不截断、不丢内容。"""
        if len(text) <= REWRITE_PART_CHARS:
            return [text]
        parts: List[str] = []
        buf = ""
        for para in re.split(r"\n\s*\n", text):
            piece = para.strip()
            if not piece:
                continue
            # 单段超长：先排空 buf 再硬切该段
            while len(piece) > REWRITE_PART_CHARS:
                if buf:
                    parts.append(buf)
                    buf = ""
                parts.append(piece[:REWRITE_PART_CHARS])
                piece = piece[REWRITE_PART_CHARS:]
            if buf and len(buf) + len(piece) + 2 > REWRITE_PART_CHARS:
                parts.append(buf)
                buf = piece
            else:
                buf = f"{buf}\n\n{piece}" if buf else piece
        if buf:
            parts.append(buf)
        return parts

    def _rewrite_part(self, part: str, titles_str: str, filename: str, need_meta: bool) -> dict:
        """改写单个分段（重试 3 次）。need_meta 时额外产出 title/summary/tags。"""
        import time

        system_prompt, user_prompt = self._rewrite_prompts(part, titles_str, filename, need_meta)

        last_error = None
        for attempt in range(3):
            try:
                raw_resp = _llm(system_prompt, user_prompt, temperature=0.3)
                m = re.search(r"\{.*\}", raw_resp, re.S)
                if m:
                    result = json.loads(m.group(0))
                    if isinstance(result, dict) and str(result.get("body") or "").strip():
                        return result
                last_error = ValueError("LLM response did not contain a usable JSON body")
            except KnowledgeImportCancelled:
                raise
            except Exception as e:
                last_error = e
                if attempt < 2:
                    wait = 3 * (attempt + 1)
                    logger.info(f"[source_to_md] LLM attempt {attempt + 1} failed ({e}), retrying in {wait}s...")
                    time.sleep(wait)

        raise RuntimeError(f"LLM optimization failed after 3 attempts: {last_error}")

    @staticmethod
    def _rewrite_prompts(part: str, titles_str: str, filename: str, need_meta: bool) -> tuple:
        task_lines = (
            "Your tasks:\n"
            "1. Generate a SHORT title of 4-6 characters in the document's language. "
            "Do NOT use the original filename. Capture the core topic concisely.\n"
            "2. Write a concise 2-3 sentence summary of the document.\n"
            "3. Rewrite the body as clean, well-organized Markdown:\n"
        ) if need_meta else (
            "Your task:\n"
            "Rewrite this body segment as clean, well-organized Markdown:\n"
        )
        json_spec = (
            "Return ONLY valid JSON with keys: title, summary, body, tags.\n"
            "- title: short title (4-6 characters)\n"
            "- summary: 2-3 sentence summary text\n"
            "- body: the cleaned markdown body (no frontmatter, no source links)\n"
            "- tags: array of tag strings\n"
        ) if need_meta else (
            "Return ONLY valid JSON with key: body.\n"
            "- body: the cleaned markdown body of this segment (no frontmatter, no source links)\n"
        )
        system_prompt = (
            "You are a knowledge management assistant. You convert raw document text "
            "into well-structured Obsidian-style Markdown.\n\n"
            "CRITICAL LANGUAGE RULE: All output (title, summary, body, tags) MUST be written in Chinese. "
            "Only keep English for proper nouns, product names, technical abbreviations (e.g. CoWoS, CPO, BOE), "
            "and terms that have no standard Chinese translation. Everything else must be Chinese.\n\n"
            + task_lines +
            "   - REMOVE all: page headers/footers, page numbers, author names/emails/affiliations, "
            "copyright/DOI/ISBN lines, conference or journal headers, acknowledgments, "
            "references/bibliography sections, appendices, and figure/image captions "
            "(e.g. lines starting with 'Fig.' or 'Figure' or 'Table' followed by descriptions).\n"
            "   - Keep all substantive technical content: abstract, introduction, methodology, "
            "results, discussion, conclusion.\n"
            "   - CRITICAL: Preserve every '<!-- IMG:... -->' marker EXACTLY where it appears, "
            "on its own line, unchanged. These markers reference embedded images; do NOT "
            "remove, merge, reorder, or modify them.\n"
            "   - Use proper heading hierarchy (##, ###).\n"
            "   - Fix broken paragraphs, merged words, and formatting artifacts from PDF extraction.\n"
            "   - Insert [[wikilinks]] to relevant existing notes where natural. Only link "
            "to notes that genuinely relate to the content. Do not force links.\n"
            "   - Do not include frontmatter or source links in the body.\n\n"
            + json_spec
        )
        user_prompt = (
            f"File: {filename}\n"
            f"Existing vault notes (insert wikilinks where relevant): {titles_str}\n\n"
            f"---RAW CONTENT---\n{part}"
        )
        return system_prompt, user_prompt

    async def _rewrite_part_async(
        self,
        part: str,
        titles_str: str,
        filename: str,
        need_meta: bool,
        on_progress: Optional[ProgressCallback] = None,
        progress_start: int = 34,
        progress_end: int = 70,
    ) -> dict:
        """Async variant so cancellation can interrupt the LLM HTTP request."""
        import asyncio

        system_prompt, user_prompt = self._rewrite_prompts(part, titles_str, filename, need_meta)
        output_chars = 0
        target_chars = max(1, int(len(part) * 1.2))

        # Progress reporting: dual-dimension fraction (chars + elapsed time)
        # + background heartbeat + final flush. This ensures the progress bar
        # advances smoothly even when the LLM response is buffered by network
        # proxies (all tokens arrive in one burst) or when output is much
        # shorter than the input text.
        import time as _time
        _last_report_ts = _time.monotonic()
        _llm_start_ts = _time.monotonic()
        _LLM_EXPECTED_S = 30.0
        _THROTTLE_S = 0.15

        def _compute_fraction() -> float:
            """Dual-dimension: max(char_fraction, time_fraction) so progress
            advances even when output_chars is tiny (LLM still thinking) or
            when the response is buffered (all tokens arrive at once)."""
            char_frac = output_chars / target_chars
            elapsed = _time.monotonic() - _llm_start_ts
            time_frac = min(1.0, elapsed / _LLM_EXPECTED_S)
            return min(1.0, max(char_frac, time_frac))

        def _report_progress() -> None:
            if on_progress is None:
                return
            fraction = _compute_fraction()
            percent = progress_start + int((progress_end - progress_start) * fraction)
            on_progress(
                "智能改写",
                min(max(percent, progress_start), progress_end),
                f"已生成 {output_chars}/{target_chars} 字",
            )

        def handle_chunk(text: str) -> None:
            nonlocal output_chars, _last_report_ts
            output_chars += len(text)
            if on_progress is None:
                return
            now = _time.monotonic()
            if now - _last_report_ts < _THROTTLE_S:
                return
            _last_report_ts = now
            _report_progress()

        async def _heartbeat() -> None:
            """Report progress every 2 seconds while the LLM call is in-flight.
            Ensures smooth progress even if the SSE response is fully buffered
            by a network proxy and handle_chunk is never called until the end."""
            while True:
                await asyncio.sleep(2.0)
                _report_progress()
                nonlocal _last_report_ts
                _last_report_ts = _time.monotonic()

        last_error = None
        for attempt in range(3):
            try:
                _llm_start_ts = _time.monotonic()
                heartbeat_task = asyncio.create_task(_heartbeat()) if on_progress else None
                try:
                    raw_resp = await _llm_async(
                        system_prompt,
                        user_prompt,
                        temperature=0.3,
                        on_chunk=handle_chunk if on_progress else None,
                    )
                finally:
                    if heartbeat_task:
                        heartbeat_task.cancel()
                        try:
                            await heartbeat_task
                        except asyncio.CancelledError:
                            pass
                # Final flush: push progress to progress_end so the bar
                # does not stall at the last throttled value.
                if on_progress:
                    on_progress("智能改写", progress_end, "本段改写完成")
                m = re.search(r"\{.*\}", raw_resp, re.S)
                if m:
                    result = json.loads(m.group(0))
                    if isinstance(result, dict) and str(result.get("body") or "").strip():
                        return result
                last_error = ValueError("LLM response did not contain a usable JSON body")
            except KnowledgeImportCancelled:
                raise
            except Exception as e:
                last_error = e
                if attempt < 2:
                    wait = 3 * (attempt + 1)
                    logger.info(f"[source_to_md] LLM attempt {attempt + 1} failed ({e}), retrying in {wait}s...")
                    await asyncio.sleep(wait)

        raise RuntimeError(f"LLM optimization failed after 3 attempts: {last_error}")

    def _llm_optimize(self, raw_md: str, existing_titles: List[str], filename: str) -> dict:
        """用 KNOWLEDGE_REWRITE_MODEL 全文改写：title/summary/tags/wikilinks/正文。

        不做字符截断：超过单段上限时按段落边界分段顺序改写再拼接，全文覆盖。
        Returns dict with keys: title, summary, body, tags.
        Raises RuntimeError if LLM fails after retries（调用方决定降级策略）。
        """
        parts = self._split_for_rewrite(raw_md)
        titles_str = ", ".join(existing_titles[:50]) if existing_titles else "(none)"
        title, summary = "", ""
        tags: List[str] = []
        bodies: List[str] = []
        for i, part in enumerate(parts):
            result = self._rewrite_part(part, titles_str, filename, need_meta=(i == 0))
            if i == 0:
                title = result.get("title", "")
                summary = result.get("summary", "")
                tags = result.get("tags", [])
            bodies.append(str(result.get("body") or part))
        return {"title": title, "summary": summary, "tags": tags, "body": "\n\n".join(bodies)}

    async def _llm_optimize_async(
        self,
        raw_md: str,
        existing_titles: List[str],
        filename: str,
        on_progress: Optional[ProgressCallback] = None,
        should_cancel: Optional[CancelCheck] = None,
    ) -> dict:
        parts = self._split_for_rewrite(raw_md)
        titles_str = ", ".join(existing_titles[:50]) if existing_titles else "(none)"
        title, summary = "", ""
        tags: List[str] = []
        bodies: List[str] = []
        for i, part in enumerate(parts):
            if should_cancel and should_cancel():
                raise KnowledgeImportCancelled()
            part_start = 34 + round(i / len(parts) * 36)
            part_end = 34 + round((i + 1) / len(parts) * 36)
            if on_progress:
                on_progress("智能改写", part_start, f"第 {i + 1}/{len(parts)} 段")
            result = await self._rewrite_part_async(
                part,
                titles_str,
                filename,
                need_meta=(i == 0),
                on_progress=on_progress,
                progress_start=part_start,
                progress_end=part_end,
            )
            if i == 0:
                title = result.get("title", "")
                summary = result.get("summary", "")
                tags = result.get("tags", [])
            bodies.append(str(result.get("body") or part))
        return {"title": title, "summary": summary, "tags": tags, "body": "\n\n".join(bodies)}

    # ---- Frontmatter and source link ----

    def _build_frontmatter(
        self, title: str, source_file: str, original_filename: str,
        tags: List[str], summary: str,
    ) -> str:
        ts = datetime.now().strftime("%Y-%m-%d")
        tag_yaml = "\n".join(f"  - {t}" for t in tags) if tags else "[]"
        clean_summary = summary.replace('"', '').replace('\n', ' ') if summary else ""
        summary_line = f'summary: "{clean_summary}"' if summary else ""

        lines = [
            "---",
            f"title: {title}",
            f"source_file: {source_file}",
            f"original_filename: {original_filename}",
            f"converted_at: {ts}",
            f"type: source",
            f"status: active",
            f"created: {ts}",
            f"updated: {ts}",
        ]
        if summary_line:
            lines.append(summary_line)
        if tags:
            lines.append("tags:")
            for t in tags:
                lines.append(f"  - {t}")
        else:
            lines.append("tags: []")
        lines.append("---\n")
        return "\n".join(lines)

    def _source_link(self, stored_name: str, ext: str, kb_id: int) -> str:
        # 文件名做 URL 编码：原始文件名可能含 LaTeX（$ { } ^ \）、空格、括号等字符，
        # 直接拼进链接会破坏 CommonMark 解析（反斜杠转义/空格截断），前端渲染不出可点击链接；
        # 下载端点的路径参数会被自动解码，编码对既有接口透明
        from urllib.parse import quote
        return f"\n\n---\n\n[View original file](/api/v1/knowledge/files/{kb_id}/{quote(stored_name, safe='')})\n"
